"""
excel_handler.py

Reads the Layer_Servers sheet and writes validation results.
"""

from pathlib import Path

from openpyxl import load_workbook

from config import INPUT_EXCEL, OUTPUT_EXCEL
from logger import logger


class ExcelHandler:

    OUTPUT_COLUMNS = [
        "Status",
        "HTTP Status",
        "Feature Count",
        "Error",
        "Screenshot",
    ]

    def __init__(self):

        # Resume from previous output if it exists
        if Path(OUTPUT_EXCEL).exists():
            self.workbook = load_workbook(OUTPUT_EXCEL)
            logger.info(f"Resuming validation from {OUTPUT_EXCEL}")
        else:
            self.workbook = load_workbook(INPUT_EXCEL)
            logger.info(f"Starting validation using {INPUT_EXCEL}")

        # Always use the Layer_Servers sheet
        self.sheet = self.workbook["Layer_Servers"]

        self.url_column = None
        self.output_columns = {}

        self._find_url_column()
        self._create_output_columns()

    def _find_url_column(self):
        """Locate the arcgis_link column."""

        for cell in self.sheet[1]:

            if cell.value is None:
                continue

            header = str(cell.value).strip().lower()

            if header == "arcgis_link":
                self.url_column = cell.column
                logger.info(
                    f"Using ArcGIS URL column: {cell.column_letter} ({cell.value})"
                )
                return

        raise Exception("Column 'arcgis_link' not found.")

    def _create_output_columns(self):
        """Create output columns if they do not already exist."""

        existing = {}

        for cell in self.sheet[1]:
            if cell.value:
                existing[str(cell.value)] = cell.column

        start_col = self.sheet.max_column + 1

        for name in self.OUTPUT_COLUMNS:

            if name in existing:
                self.output_columns[name] = existing[name]

            else:
                self.sheet.cell(row=1, column=start_col).value = name
                self.output_columns[name] = start_col
                start_col += 1

    def iter_rows(self):
        """Yield only rows that have not yet been validated."""

        status_col = self.output_columns["Status"]

        for row in range(2, self.sheet.max_row + 1):

            # Skip rows already processed
            status = self.sheet.cell(
                row=row,
                column=status_col,
            ).value

            if status is not None and str(status).strip():
                continue

            url = self.sheet.cell(
                row=row,
                column=self.url_column,
            ).value

            if url is None:
                continue

            url = str(url).strip()

            if not url:
                continue

            yield row, url

    def write_result(self, row, result):

        self.sheet.cell(
            row=row,
            column=self.output_columns["Status"],
        ).value = result.status

        self.sheet.cell(
            row=row,
            column=self.output_columns["HTTP Status"],
        ).value = result.http_status

        self.sheet.cell(
            row=row,
            column=self.output_columns["Feature Count"],
        ).value = result.feature_count

        self.sheet.cell(
            row=row,
            column=self.output_columns["Error"],
        ).value = result.error

        self.sheet.cell(
            row=row,
            column=self.output_columns["Screenshot"],
        ).value = result.screenshot

    def save(self):
        self.workbook.save(OUTPUT_EXCEL)
        logger.info(f"Saved workbook -> {OUTPUT_EXCEL}")