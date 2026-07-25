from datetime import datetime

from openpyxl import load_workbook
from sqlalchemy import (
    Column,
    DateTime,
    Integer,
    String,
    create_engine,
)
from sqlalchemy.orm import declarative_base, sessionmaker

DATABASE_URL = "sqlite:///arcgis_validation.db"

engine = create_engine(DATABASE_URL)
Session = sessionmaker(bind=engine)
Base = declarative_base()


class ValidatedLayer(Base):
    __tablename__ = "validated_layers"

    id = Column(Integer, primary_key=True, autoincrement=True)

    layer_id = Column(Integer)
    layer_key = Column(String)
    server = Column(String)
    source_sheet = Column(String)

    link = Column(String)
    arcgis_link = Column(String)

    status = Column(String)
    http_status = Column(Integer)
    feature_count = Column(Integer)

    error = Column(String)
    screenshot = Column(String)

    validated_at = Column(DateTime, default=datetime.utcnow)


Base.metadata.create_all(engine)


def load_excel_to_database():

    workbook = load_workbook("output/validated.xlsx", data_only=True)
    sheet = workbook["Layer_Servers"]

    # Read headers dynamically
    headers = {}
    for cell in sheet[1]:
        headers[cell.value] = cell.column

    session = Session()

    inserted = 0

    for row in range(2, sheet.max_row + 1):

        def value(name):
            column = headers.get(name)
            if column is None:
                return None
            return sheet.cell(row=row, column=column).value

        record = ValidatedLayer(
            layer_id=value("layer_id"),
            layer_key=value("layer_key"),
            server=value("server"),
            source_sheet=value("source_sheet"),
            link=value("link"),
            arcgis_link=value("arcgis_link"),
            status=value("Status"),
            http_status=value("HTTP Status"),
            feature_count=value("Feature Count"),
            error=value("Error"),
            screenshot=value("Screenshot"),
        )

        session.add(record)
        inserted += 1

    session.commit()
    session.close()

    print(f"Inserted {inserted} records into SQLite database.")


if __name__ == "__main__":
    load_excel_to_database()